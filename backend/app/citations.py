"""Citation formatting: APA / MLA / Chicago + inline templates. Pure fns, no network."""
from __future__ import annotations
import re
from datetime import date

STYLES = ("apa", "mla", "chicago")


def _author_last(title: str) -> str:
    # Fallback: org/site name from title when no author given
    t = (title or "").strip()
    return t[:60] if t else "n.d."


def format_inline(source_title: str = "", url: str = "", locator: str = "",
                  style: str = "apa", year: str = "") -> str:
    """Short inline citation, e.g. APA: (Title, p. 3)."""
    style = (style or "apa").lower()
    t = (source_title or url or "source").strip()
    loc = f", {locator.strip()}" if locator and locator.strip() else ""
    y = year.strip() if year else ""
    if style == "mla":
        return f'("{t}"{loc})'
    if style == "chicago":
        return f'({t}{loc}, {url})'.replace(", )", ")") if url else f"({t}{loc})"
    # apa default
    return f"({t}{f', {y}' if y else ''}{loc})"


def format_bibliography(source_title: str = "", url: str = "", locator: str = "",
                        style: str = "apa", author: str = "", year: str = "",
                        accessed: str = "") -> str:
    """Full bibliography entry."""
    style = (style or "apa").lower()
    t = (source_title or "Untitled source").strip()
    u = (url or "").strip()
    acc = accessed or date.today().isoformat()
    y = year.strip() if year else "n.d."
    a = author.strip() or _author_last(t)
    if style == "mla":
        # Author. "Title." URL, Accessed day.
        return f'{a}. "{t}." {u}, Accessed {acc}.'.strip()
    if style == "chicago":
        # Author, "Title," URL (accessed date).
        return f'{a}, "{t}," {u} (accessed {acc}).'.strip()
    # apa: Author. (Year). Title. URL
    return f"{a}. ({y}). {t}. {u}".strip()


CITATION_TEMPLATES = {
    "support": "According to {inline}, {claim} — {excerpt}",
    "contrast": "However, {inline} qualifies this: {excerpt}",
    "stat": "{claim} ({stat} per {inline})",
    "define": "{term} is defined as {excerpt} {inline}",
}


def apply_template(name: str, inline: str = "", claim: str = "",
                   excerpt: str = "", stat: str = "", term: str = "") -> str:
    tpl = CITATION_TEMPLATES.get(name, CITATION_TEMPLATES["support"])
    return tpl.format(inline=inline, claim=claim, excerpt=excerpt, stat=stat, term=term)


def extract_doc_text(filename: str, content: bytes) -> str:
    """Best-effort local text extraction for docx/csv/txt/md. xlsx needs openpyxl (optional)."""
    name = (filename or "").lower()
    try:
        if name.endswith(".docx"):
            import zipfile
            import xml.etree.ElementTree as ET
            with zipfile.ZipFile(__import__("io").BytesIO(content)) as z:
                xml = z.read("word/document.xml")
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            paras = []
            for p in ET.fromstring(xml).findall(".//w:p", ns):
                paras.append("".join(t.text or "" for t in p.findall(".//w:t", ns)))
            return "\n".join(paras)[:20000]
        if name.endswith((".csv", ".tsv", ".txt", ".md", ".markdown")):
            return content.decode("utf-8", errors="replace")[:20000]
        if name.endswith(".xlsx"):
            try:
                from openpyxl import load_workbook  # optional dep
                import io
                wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
                out = []
                for ws in wb.worksheets[:4]:
                    out.append(f"## Sheet: {ws.title}")
                    for row in list(ws.iter_rows(values_only=True))[:60]:
                        out.append(", ".join("" if v is None else str(v) for v in row))
                return "\n".join(out)[:20000]
            except ImportError:
                return ""
    except Exception:
        return ""
    return ""
